"""
Atualização da proposta comercial em PowerPoint.

O modelo (modelo/Proposta_Comercial.pptx) usa marcadores no formato {{NOME}}
em caixas de texto, formas, tabelas e grupos. Ao gerar, cada marcador é trocado
pelo valor da proposta mantendo a formatação do texto original.

O PPT oficial pode ser usado do jeito que está: os textos de exemplo destacados
nele (ex.: "INSERIR RAZÃO SOCIAL", a linha "Compensação ... ano base ...", a
forma de pagamento e o contato do último slide) são reconhecidos e trocados
pelos marcadores automaticamente (TEXTOS_PADRAO).
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import math
import re

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches

from .calculos import fmt_data, fmt_moeda, fmt_num, fmt_pct, juntar, texto_escopo
from .ooxml_strict import converter, eh_strict
from .ppt_tabela import inserir_tabela

BASE_DIR = Path(__file__).resolve().parent.parent
PPT_MODEL_PATH = BASE_DIR / "modelo" / "Proposta_Comercial.pptx"

PATTERN = re.compile(r"\{\{\s*([A-Z0-9_]+)\s*\}\}")

# Descrição exibida no painel e no README. Marcadores por ano (META_2025,
# RECUPERACAO_2025, VOLUME_2025, COMPENSAR_2025, VALOR_2025) são gerados para cada ano.
DESCRICOES = {
    "CLIENTE": "Razão social do cliente",
    "CNPJ": "CNPJ do cliente",
    "CONTATO": "Nome do contato no cliente",
    "EMAIL_CLIENTE": "E-mail do contato",
    "TELEFONE": "Telefone do contato",
    "CIDADE_UF": "Cidade/UF do cliente",
    "NUMERO_PROPOSTA": "Número da proposta",
    "DATA_PROPOSTA": "Data da proposta (dd/mm/aaaa)",
    "VALIDADE": "Data de validade (dd/mm/aaaa)",
    "VALIDADE_DIAS": "Validade em dias",
    "VALIDADE_EXTENSO": "Validade por extenso, ex.: 30 (trinta)",
    "RESPONSAVEL": "Responsável comercial",
    "EMAIL_RESPONSAVEL": "E-mail do responsável comercial",
    "TELEFONE_RESPONSAVEL": "Telefone do responsável comercial",
    "FORMA_PAGAMENTO": "Forma de pagamento, ex.: 30 dias após o aceite desta proposta.",
    "PRECO_TONELADA": "Preço por tonelada compensada (R$)",
    "PLANO": "Plano até 5 t (ex.: Plano até 3 t) ou cálculo por tonelada",
    "ESCOPO": "Ex.: Compensação SP e MS ano base 2019 e 2020 e nacional ano base 2024 e 2025",
    "TEXTO_METAS": "Texto legal das metas mínimas dos anos da proposta",
    "RECUPERACAO": "Recuperação contratada por ano (ex.: 2025: 40%; 2024: 32%)",
    "ANOS": "Anos de compensação (ex.: 2025, 2024)",
    "QTD_ANOS": "Quantidade de anos",
    "ESTADOS": "UFs com volume declarado",
    "QTD_ESTADOS": "Quantidade de UFs",
    "MATERIAIS": "Materiais com volume declarado",
    "VOLUME_DECLARADO": "Massa total declarada (t)",
    "VOLUME_COMPENSAR": "Massa total a compensar (t)",
    "SUBTOTAL": "Valor antes do desconto (R$)",
    "DESCONTO_PERCENTUAL": "Desconto (%)",
    "DESCONTO": "Desconto (R$)",
    "VALOR_TOTAL": "Valor total da proposta (R$)",
    "OBSERVACOES": "Observações da proposta",
    "EMPRESA": "Razão social da Ambipar",
    "CNPJ_EMPRESA": "CNPJ da Ambipar",
}


UNIDADES = ["zero", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove", "dez",
            "onze", "doze", "treze", "quatorze", "quinze", "dezesseis", "dezessete", "dezoito", "dezenove"]
DEZENAS = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta", "oitenta", "noventa"]
CENTENAS = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos", "seiscentos",
            "setecentos", "oitocentos", "novecentos"]


def numero_extenso(n):
    """Inteiros de 0 a 999 por extenso (suficiente para dias de validade)."""
    if n < 20:
        return UNIDADES[n]
    if n < 100:
        d, u = divmod(n, 10)
        return DEZENAS[d] + (f" e {UNIDADES[u]}" if u else "")
    if n == 100:
        return "cem"
    if n < 1000:
        c, resto = divmod(n, 100)
        return CENTENAS[c] + (f" e {numero_extenso(resto)}" if resto else "")
    return str(n)


def montar_valores(proposta, calc, parametros, ctx):
    cliente, dados = proposta["cliente"], proposta["proposta"]
    contato = parametros.get("contato_padrao", {"email": "", "telefone": ""})
    por_ano = {a["ano"]: a for a in calc["por_ano"]}
    # Texto legal do ano mais recente + resumo curto dos demais (cabe no slide).
    recente = max(calc["anos"])
    regra = por_ano[recente]["regra"]
    textos = [regra["texto"] or f'Ano-base {recente}: {regra["resumo"]}']
    demais = []
    for a in sorted(calc["anos"]):
        if a == recente:
            continue
        r = por_ano[a]["regra"]
        exc = "".join(f", {uf} {fmt_pct(p / 100)}" for uf, p in r["excecoes"].items())
        demais.append(f'{a} ({fmt_pct(r["minimo"] / 100)}{exc})')
    if demais:
        textos.append(f"Demais anos-base: {juntar(demais)}.")

    valores = {
        "CLIENTE": cliente["razao_social"],
        "CNPJ": cliente["cnpj"],
        "CONTATO": cliente["contato"],
        "EMAIL_CLIENTE": cliente["email"],
        "TELEFONE": cliente["telefone"],
        "CIDADE_UF": cliente["cidade_uf"],
        "NUMERO_PROPOSTA": dados["numero"],
        "DATA_PROPOSTA": fmt_data(dados["data"]),
        "VALIDADE": fmt_data(dados["validade"]),
        "VALIDADE_DIAS": str(dados["validade_dias"]),
        "VALIDADE_EXTENSO": f'{dados["validade_dias"]} ({numero_extenso(dados["validade_dias"])})',
        "RESPONSAVEL": dados["responsavel"],
        "EMAIL_RESPONSAVEL": dados["email_responsavel"] or contato["email"],
        "TELEFONE_RESPONSAVEL": dados["telefone_responsavel"] or contato["telefone"],
        "FORMA_PAGAMENTO": dados["forma_pagamento"],
        "PRECO_TONELADA": fmt_moeda(calc["preco_efetivo"]),
        "PLANO": (f'Plano até {fmt_num(calc["plano"]["ate_t"], 0)} t' if calc["regime"] == "plano"
                  else f'{fmt_num(calc["compensar"], 3)} t a compensar'),
        "ESCOPO": texto_escopo(calc),
        "TEXTO_METAS": " ".join(textos),
        "RECUPERACAO": "; ".join(
            f'{a["ano"]}: {fmt_pct(a["recuperacao"])}' for a in calc["por_ano"]),
        "ANOS": juntar(sorted(calc["anos"])),
        "QTD_ANOS": str(len(calc["anos"])),
        "ESTADOS": juntar(e["uf"] for e in calc["por_estado"]),
        "QTD_ESTADOS": str(len(calc["por_estado"])),
        "MATERIAIS": juntar(m["material"].title() for m in calc["por_material"]),
        "VOLUME_DECLARADO": f'{fmt_num(calc["declarado"], 3)} t',
        "VOLUME_COMPENSAR": f'{fmt_num(calc["compensar"], 3)} t',
        "SUBTOTAL": fmt_moeda(calc["subtotal"]),
        "DESCONTO_PERCENTUAL": fmt_pct(calc["desconto_percentual"]),
        "DESCONTO": fmt_moeda(calc["desconto_valor"]),
        "VALOR_TOTAL": fmt_moeda(calc["total"]),
        "OBSERVACOES": dados["observacoes"],
        "EMPRESA": parametros["empresa"]["razao_social"],
        "CNPJ_EMPRESA": parametros["empresa"]["cnpj"],
    }

    for ano in ctx["anos"]:
        info = por_ano.get(ano)
        valores[f"META_{ano}"] = ctx["regras"][ano]["resumo"]
        valores[f"RECUPERACAO_{ano}"] = fmt_pct(info["recuperacao"]) if info else "-"
        valores[f"VOLUME_{ano}"] = f'{fmt_num(info["declarado"], 3)} t' if info else "-"
        valores[f"COMPENSAR_{ano}"] = f'{fmt_num(info["compensar"], 3)} t' if info else "-"
        valores[f"VALOR_{ano}"] = fmt_moeda(info["valor"]) if info else "-"

    return {k: str(v).replace("\r", "").replace("\n", " ") for k, v in valores.items()}


def listar_marcadores(anos_modelo):
    lista = [{"marcador": f"{{{{{k}}}}}", "descricao": v} for k, v in DESCRICOES.items()]
    if anos_modelo:
        ano = anos_modelo[0]
        lista += [
            {"marcador": f"{{{{META_{ano}}}}}", "descricao": "Regra mínima do ano (troque o ano)"},
            {"marcador": f"{{{{RECUPERACAO_{ano}}}}}", "descricao": "Recuperação contratada no ano"},
            {"marcador": f"{{{{VOLUME_{ano}}}}}", "descricao": "Massa declarada no ano (t)"},
            {"marcador": f"{{{{COMPENSAR_{ano}}}}}", "descricao": "Massa a compensar no ano (t)"},
            {"marcador": f"{{{{VALOR_{ano}}}}}", "descricao": "Valor do ano (R$)"},
        ]
    return lista


def _replace_paragraph(paragraph, valores, pendentes):
    runs = paragraph.runs
    if not runs:
        return

    def sub(match):
        key = match.group(1)
        if key in valores:
            return valores[key]
        pendentes.add(key)
        return match.group(0)

    # 1) Marcador inteiro dentro de um único trecho: troca preservando cada formatação.
    for run in runs:
        if "{{" in run.text:
            run.text = PATTERN.sub(sub, run.text)

    # 2) PowerPoint às vezes quebra "{{CLIENTE}}" em vários trechos; junta no primeiro.
    full = "".join(run.text for run in runs)
    if any(m.group(1) in valores for m in PATTERN.finditer(full)):
        runs[0].text = PATTERN.sub(sub, full)
        for run in runs[1:]:
            run.text = ""


# (regex sobre o texto do parágrafo, substituição). Só valem para o slide de contato
# os itens marcados com True.
TEXTOS_PADRAO = [
    (re.compile(r"INSERIR\s+RAZ[ÃA]O\s+SOCIAL", re.I), "{{CLIENTE}}", False),
    (re.compile(r"^\s*Compensa[çc][ãa]o\b.*\bano\s+base\b.*$", re.I | re.S), "{{ESCOPO}}", False),
    (re.compile(r"^\s*A meta m[íi]nima para a comprova[çc][ãa]o da log[íi]stica reversa.*$", re.I | re.S),
     "{{TEXTO_METAS}}", False),
    (re.compile(r"30\s*dias\s+ap[óo]s\s+o\s+aceite\s+desta\s+proposta\.?", re.I), "{{FORMA_PAGAMENTO}}", False),
    (re.compile(r"v[áa]lida\s+por\s+30\s*\(trinta\)\s*dias", re.I), "válida por {{VALIDADE_EXTENSO}} dias", False),
    (re.compile(r"[\w.]+@ambipar\.co\s*m\b", re.I), "{{EMAIL_RESPONSAVEL}}", True),
    (re.compile(r"\+?55\s*\(?11\)?\s*9?\d{4}-?\d{4}"), "{{TELEFONE_RESPONSAVEL}}", True),
]


def _slide_de_contato(slide):
    texto = " ".join(tf.text for tf in _iter_text_frames(slide.shapes)).upper()
    return "OBRIGADO" in texto or "E N T R E" in texto or "ENTRE EM CONTATO" in texto


def _encolher_se_cresceu(paragraph, tamanho_original):
    """Texto longo do modelo que cresceu (ex.: metas de vários anos): reduz a fonte até 70%."""
    novo = len(paragraph.text)
    if tamanho_original < 60 or novo <= tamanho_original * 1.08:
        return
    fator = max(0.7, tamanho_original / novo)
    for run in paragraph.runs:
        if run.font.size:
            run.font.size = int(run.font.size * fator)


def _trocar_trecho(runs, inicio, fim, marcador):
    """
    Troca o texto entre as posições [inicio, fim) do parágrafo pelo marcador,
    mexendo só nos trechos (runs) atingidos: a formatação do resto fica igual.
    """
    pos = 0
    limites = []
    for run in runs:
        limites.append((pos, pos + len(run.text)))
        pos += len(run.text)
    i = next(k for k, (a, b) in enumerate(limites) if a <= inicio < b or (a == b == inicio))
    j = next(k for k, (a, b) in enumerate(limites) if a < fim <= b)
    cabeca = runs[i].text[:inicio - limites[i][0]]
    cauda = runs[j].text[fim - limites[j][0]:]
    if i == j:
        runs[i].text = cabeca + marcador + cauda
        return
    runs[i].text = cabeca + marcador
    for k in range(i + 1, j):
        runs[k].text = ""
    runs[j].text = cauda


def _marcar_textos_padrao(paragraph, contato):
    """Troca os textos de exemplo do PPT oficial por marcadores. Devolve True se trocou."""
    runs = paragraph.runs
    if not runs:
        return False
    mudou = False
    for regex, marcador, so_contato in TEXTOS_PADRAO:
        if so_contato and not contato:
            continue
        for _ in range(10):  # mais de uma ocorrência no mesmo parágrafo
            texto = "".join(run.text for run in runs)
            achado = regex.search(texto)
            if not achado or achado.group(0) == marcador:
                break
            _trocar_trecho(runs, achado.start(), achado.end(), marcador)
            mudou = True
    return mudou


def _tirar_destaques(paragraph):
    """Tira o destaque amarelo de "preencher aqui" usado no modelo."""
    for destaque in paragraph._p.findall(f'.//{qn("a:highlight")}'):
        destaque.getparent().remove(destaque)


def _iter_text_frames(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_text_frames(shape.shapes)
            continue
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


FORMA_PAGAMENTO = re.compile(r"forma\s+de\s+pagamento", re.I)
LINHA_OBS = Inches(0.21)          # altura aproximada de uma linha de 11 pt com espaçamento
CARACTERES_POR_POLEGADA = 13      # ~11 pt Neo Sans


def _inserir_observacoes(prs, observacoes):
    """
    Coloca as observações da proposta logo abaixo da linha "Forma de pagamento",
    com a mesma formatação. A caixa de texto do modelo cresce para baixo e já
    encosta no rodapé, então ela sobe o necessário, sem passar da tabela acima.
    """
    linhas = [linha.strip() for linha in observacoes.splitlines() if linha.strip()]
    if not linhas:
        return False
    for slide in prs.slides:
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False) or not shape.has_text_frame:
                continue
            paragrafos = shape.text_frame.paragraphs
            alvo = next((p for p in paragrafos if FORMA_PAGAMENTO.search(p.text)), None)
            if alvo is None or not alvo.runs:
                continue

            anterior = alvo._p
            for k, linha in enumerate(linhas):
                novo = deepcopy(alvo._p)
                runs = novo.findall(qn("a:r"))
                for extra in runs[1:]:
                    novo.remove(extra)
                for quebra in novo.findall(qn("a:br")):
                    novo.remove(quebra)
                runs[0].find(qn("a:t")).text = (f"*Observações: {linha}" if k == 0 else linha)
                anterior.addnext(novo)
                anterior = novo

            # Sobe a caixa o quanto as linhas novas ocupam, sem invadir a tabela.
            largura = Emu(shape.width).inches
            n_linhas = sum(math.ceil(len(l) / max(largura * CARACTERES_POR_POLEGADA, 1)) for l in linhas)
            limite = max((s.top + s.height for s in slide.shapes
                          if getattr(s, "has_table", False) and s.has_table and s.top < shape.top),
                         default=0) + Inches(0.15)
            subir = min(n_linhas * LINHA_OBS, max(shape.top - limite, 0))
            shape.top = shape.top - int(subir)
            return True
    return False


def gerar_ppt(valores, tabela=None, observacoes=""):
    if not PPT_MODEL_PATH.exists():
        raise FileNotFoundError(
            "Modelo da proposta não encontrado. Salve o PowerPoint com os marcadores "
            f"{{{{...}}}} em: modelo/{PPT_MODEL_PATH.name}"
        )

    conteudo = PPT_MODEL_PATH.read_bytes()
    if eh_strict(conteudo):
        # Modelo salvo como "Open XML Strict": converte em memória para o formato padrão.
        conteudo = converter(conteudo)
    prs = Presentation(BytesIO(conteudo))
    if tabela:
        inserir_tabela(prs, tabela)
    # Se o modelo já tiver o marcador {{OBSERVACOES}}, as observações vão só para ele.
    modelo_texto = " ".join(tf.text for slide in prs.slides for tf in _iter_text_frames(slide.shapes))
    pendentes = set()
    for slide in prs.slides:
        contato = _slide_de_contato(slide)
        for text_frame in _iter_text_frames(slide.shapes):
            for paragraph in text_frame.paragraphs:
                tamanho_original = len(paragraph.text)
                marcado = _marcar_textos_padrao(paragraph, contato)
                _replace_paragraph(paragraph, valores, pendentes)
                if marcado:
                    _encolher_se_cresceu(paragraph, tamanho_original)
                _tirar_destaques(paragraph)

    if observacoes and "{{OBSERVACOES}}" not in modelo_texto:
        _inserir_observacoes(prs, observacoes)

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output, sorted(pendentes)
