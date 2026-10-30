"""
Planilha de autodeclaração.

O arquivo modelo/Declaracao_volume_embalagens.xlsx é tratado como template
intocável: abrimos o modelo, preenchemos somente B5:K31 e devolvemos uma cópia.
Ano-base sem guia no modelo (ex.: 2026) ganha, só na cópia, uma guia igual à
do ano mais recente, com o ano em B1 e a meta do ano em L1.
"""
from io import BytesIO
from pathlib import Path
import re

from openpyxl import load_workbook

from .calculos import ESTADOS, MATERIAIS, normalize_number

BASE_DIR = Path(__file__).resolve().parent.parent
MODEL_PATH = BASE_DIR / "modelo" / "Declaracao_volume_embalagens.xlsx"

FIRST_STATE_ROW = 5
LAST_STATE_ROW = 31
FIRST_MATERIAL_COL = 2
LAST_MATERIAL_COL = 11
META_CELL = "L1"


def load_model():
    if not MODEL_PATH.exists():
        raise FileNotFoundError(f"Modelo não encontrado: {MODEL_PATH}")
    return load_workbook(MODEL_PATH, data_only=False, read_only=False)


def _year_sheets(wb):
    return [ws for ws in wb.worksheets if re.fullmatch(r"\d{4}", str(ws.title).strip())]


def get_available_years():
    return sorted((int(ws.title) for ws in _year_sheets(load_model())), reverse=True)


def get_metas():
    """Meta de compensação de cada ano (célula L1 de cada guia), ex.: {2025: 0.22}."""
    metas = {}
    for ws in _year_sheets(load_model()):
        try:
            metas[int(ws.title)] = normalize_number(ws[META_CELL].value)
        except ValueError:
            metas[int(ws.title)] = 0.0
    return metas


def get_state_rows(ws):
    rows = {}
    for row in range(FIRST_STATE_ROW, LAST_STATE_ROW + 1):
        value = ws.cell(row=row, column=1).value
        if value is not None:
            state = str(value).strip().upper()
            if re.fullmatch(r"[A-Z]{2}", state):
                rows[state] = row
    return rows


def _criar_guias(wb, anos, metas):
    guias = _year_sheets(wb)
    if not guias:
        raise ValueError("A planilha modelo não tem nenhuma guia de ano para servir de base.")
    base = max(guias, key=lambda ws: int(ws.title))
    for ano in anos:
        if str(ano) in wb.sheetnames:
            continue
        ws = wb.copy_worksheet(base)
        ws.title = str(ano)
        ws["B1"] = ano
        if ano in metas:
            ws[META_CELL] = metas[ano]
    wb._sheets.sort(key=lambda ws: -int(ws.title) if re.fullmatch(r"\d{4}", ws.title) else 0)


def generate_excel(volumes, metas=None):
    """
    Todos os estados ficam com zero em todas as guias; os informados
    sobrescrevem os zeros. Não recria nem formata a planilha.
    """
    wb = load_model()
    _criar_guias(wb, volumes.keys(), metas or {})

    for ws in wb.worksheets:
        state_rows = get_state_rows(ws)
        for state in ESTADOS:
            if state not in state_rows:
                raise ValueError(f'O estado "{state}" não foi encontrado na guia "{ws.title}".')
            for col in range(FIRST_MATERIAL_COL, LAST_MATERIAL_COL + 1):
                ws.cell(row=state_rows[state], column=col).value = 0

    for year, states_data in volumes.items():
        ws = wb[str(year)]
        state_rows = get_state_rows(ws)
        for state, values in states_data.items():
            for offset, value in enumerate(values):
                ws.cell(row=state_rows[state], column=FIRST_MATERIAL_COL + offset).value = value

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def import_excel(stream):
    """
    Lê uma planilha no mesmo layout do modelo (guias por ano, estados em A5:A31,
    materiais em B:K) e devolve {"2025": {"SP": [...]}} apenas com estados preenchidos.
    """
    try:
        wb = load_workbook(stream, data_only=True, read_only=False)
    except Exception:
        raise ValueError("Não foi possível ler o arquivo. Envie um .xlsx no layout do modelo.")

    anos_modelo = set(get_available_years())
    result = {}
    for ws in _year_sheets(wb):
        year = int(ws.title)
        if year not in anos_modelo:
            continue
        for state, row in get_state_rows(ws).items():
            if state not in ESTADOS:
                continue
            vals = []
            for col in range(FIRST_MATERIAL_COL, LAST_MATERIAL_COL + 1):
                try:
                    vals.append(normalize_number(ws.cell(row=row, column=col).value))
                except ValueError:
                    vals.append(0.0)
            if any(v > 0 for v in vals):
                result.setdefault(str(year), {})[state] = vals

    if not result:
        raise ValueError(
            "Nenhum volume encontrado. O arquivo precisa ter guias por ano "
            f"(ex.: 2025) com estados na coluna A e os {len(MATERIAIS)} materiais em B:K."
        )
    return result
