"""
Registro das propostas feitas (banco de dados local).

Cada proposta fica em dados/propostas/<numero>.json com o que foi preenchido no
painel (para reabrir), o resumo calculado e os dados da negociação (status,
valor negociado, data da venda). A planilha dados/registro_propostas.xlsx é
regenerada a partir desses arquivos sempre que algo muda, para relatórios no Excel.
"""
from datetime import date, datetime
from io import BytesIO
from pathlib import Path
import json
import threading

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .calculos import normalize_number, slug

BASE_DIR = Path(__file__).resolve().parent.parent
DADOS_DIR = BASE_DIR / "dados"
PROPOSTAS_DIR = DADOS_DIR / "propostas"
PLANILHA_PATH = DADOS_DIR / "registro_propostas.xlsx"

STATUS = ["Em elaboração", "Enviada", "Em negociação", "Ganha", "Perdida"]
STATUS_INICIAL = "Enviada"

_lock = threading.Lock()


def _agora():
    return datetime.now().isoformat(timespec="seconds")


def _arquivo(numero):
    return PROPOSTAS_DIR / f"{slug(numero, 60)}.json"


def _ler(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _gravar(path, dados):
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp")
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(dados, f, ensure_ascii=False, indent=2)
    tmp.replace(path)


def _resumo(proposta, calc):
    cliente, dados = proposta["cliente"], proposta["proposta"]
    detalhes = []
    regra_preco = (f'Plano até {calc["plano"]["ate_t"]:g} t' if calc["regime"] == "plano"
                   else "DRE (acima de 5 t)")
    for ano in calc["por_ano"]:
        for m in ano["materiais"]:
            detalhes.append({
                "ano": ano["ano"], "modo": "Nacional (CONFAZ)" if ano["modo"] == "nacional" else "Por estados",
                "recuperacao": ano["recuperacao"], "material": m["material"].title(),
                "categoria": m["categoria"].title(), "declarado": m["declarado"],
                "compensar": m["compensar"], "valor": m["valor"], "regra_preco": regra_preco,
            })
    return {
        "cliente": cliente["razao_social"], "cnpj": cliente["cnpj"], "cidade_uf": cliente["cidade_uf"],
        "contato": cliente["contato"], "email": cliente["email"], "telefone": cliente["telefone"],
        "responsavel": dados["responsavel"], "email_responsavel": dados["email_responsavel"],
        "data_proposta": dados["data"].isoformat(), "validade": dados["validade"].isoformat(),
        "anos": calc["anos"],
        "modos": sorted({"Nacional (CONFAZ)" if a["modo"] == "nacional" else "Por estados"
                         for a in calc["por_ano"]}),
        "ufs": [e["uf"] for e in calc["por_estado"]],
        "declarado": calc["declarado"], "compensar": calc["compensar"],
        "subtotal": calc["subtotal"], "desconto_percentual": calc["desconto_percentual"],
        "total": calc["total"], "regra_preco": regra_preco, "preco_t": calc["preco_efetivo"],
        "classe_mo": calc["dre"]["classe"], "custo_direto": calc["dre"]["custo_direto"],
        "resultado": calc["dre"]["resultado"], "margem": calc["dre"]["margem"], "detalhes": detalhes,
    }


def salvar(payload, proposta, calc):
    """Cria ou atualiza a proposta pelo número. Mantém status e dados da negociação."""
    numero = proposta["proposta"]["numero"]
    with _lock:
        path = _arquivo(numero)
        atual = _ler(path) if path.exists() else {}
        registro = atual.get("registro") or {
            "status": STATUS_INICIAL, "valor_negociado": None, "data_venda": "", "observacao": "",
        }
        dados = {
            "numero": numero,
            "criada_em": atual.get("criada_em") or _agora(),
            "atualizada_em": _agora(),
            "payload": payload,
            "resumo": _resumo(proposta, calc),
            "registro": registro,
        }
        _gravar(path, dados)
        return dados, _atualizar_planilha()


def _todas():
    if not PROPOSTAS_DIR.exists():
        return []
    itens = []
    for path in PROPOSTAS_DIR.glob("*.json"):
        try:
            itens.append(_ler(path))
        except (OSError, ValueError):
            continue
    return sorted(itens, key=lambda d: d.get("atualizada_em", ""), reverse=True)


def listar():
    """Lista para o painel (sem o payload completo)."""
    return [{
        **{k: v for k, v in d.items() if k != "payload"},
        "resumo": {k: v for k, v in d["resumo"].items() if k != "detalhes"},
    } for d in _todas()]


def obter(numero):
    path = _arquivo(numero)
    if not path.exists():
        raise FileNotFoundError(f"Proposta {numero} não encontrada no registro.")
    return _ler(path)


def atualizar_registro(numero, campos):
    """Atualiza status, valor negociado, data da venda e observação da negociação."""
    with _lock:
        dados = obter(numero)
        registro = dados["registro"]
        if "status" in campos:
            if campos["status"] not in STATUS:
                raise ValueError(f'Status inválido: {campos["status"]}.')
            registro["status"] = campos["status"]
        if "valor_negociado" in campos:
            bruto = campos["valor_negociado"]
            registro["valor_negociado"] = (None if bruto in (None, "")
                                           else normalize_number(bruto, "Valor negociado"))
        if "data_venda" in campos:
            bruto = str(campos["data_venda"] or "").strip()
            if bruto:
                try:
                    datetime.strptime(bruto, "%Y-%m-%d")
                except ValueError:
                    raise ValueError(f"Data da venda inválida: {bruto}")
            registro["data_venda"] = bruto
        if "observacao" in campos:
            registro["observacao"] = str(campos["observacao"] or "").strip()[:1000]
        dados["atualizada_em"] = _agora()
        _gravar(_arquivo(numero), dados)
        return dados, _atualizar_planilha()


def excluir(numero):
    with _lock:
        path = _arquivo(numero)
        if not path.exists():
            raise FileNotFoundError(f"Proposta {numero} não encontrada no registro.")
        path.unlink()
        return _atualizar_planilha()


# ---------- planilha / relatório ----------

VERDE = PatternFill("solid", fgColor="163300")
CABECALHO = Font(bold=True, color="FFFFFF")
MOEDA = '"R$" #,##0.00'
TON = "#,##0.000"
PCT = "0.00%"
DATA = "DD/MM/YYYY"


def _data(texto):
    try:
        return date.fromisoformat(texto[:10]) if texto else None
    except ValueError:
        return None


def _aba(ws, colunas, linhas):
    ws.append([c[0] for c in colunas])
    for cell in ws[1]:
        cell.fill, cell.font = VERDE, CABECALHO
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for linha in linhas:
        ws.append(linha)
    for idx, (_, largura, formato) in enumerate(colunas, start=1):
        letra = get_column_letter(idx)
        ws.column_dimensions[letra].width = largura
        if formato:
            for (cell,) in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                cell.number_format = formato
    ws.freeze_panes = "B2"
    if linhas:
        ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 30


def gerar_relatorio():
    propostas = _todas()
    wb = Workbook()

    ws = wb.active
    ws.title = "Propostas"
    colunas = [
        ("Nº da proposta", 22, None), ("Data da proposta", 13, DATA), ("Validade", 12, DATA),
        ("Cliente", 34, None), ("CNPJ", 20, None), ("Cidade/UF", 18, None), ("Contato", 22, None),
        ("E-mail", 26, None), ("Telefone", 16, None), ("Responsável", 20, None),
        ("Anos-base", 16, None), ("Modo", 22, None), ("UFs", 30, None),
        ("Massa declarada (t)", 14, TON), ("Massa a compensar (t)", 14, TON),
        ("Regra de preço", 18, None), ("Preço por t (R$)", 14, MOEDA), ("Classe MO", 9, None),
        ("Custo direto (R$)", 16, MOEDA), ("Resultado DRE (R$)", 16, MOEDA), ("Margem DRE", 11, PCT),
        ("Valor calculado (R$)", 16, MOEDA), ("Desconto", 10, PCT), ("Valor da proposta (R$)", 16, MOEDA),
        ("Valor negociado (R$)", 16, MOEDA), ("Valor final (R$)", 16, MOEDA), ("Status", 15, None),
        ("Data da venda", 13, DATA), ("Observação da negociação", 34, None),
        ("Criada em", 18, None), ("Atualizada em", 18, None),
    ]
    linhas = []
    for d in propostas:
        r, reg = d["resumo"], d["registro"]
        negociado = reg.get("valor_negociado")
        linhas.append([
            d["numero"], _data(r["data_proposta"]), _data(r["validade"]), r["cliente"], r["cnpj"],
            r["cidade_uf"], r["contato"], r["email"], r["telefone"], r["responsavel"],
            ", ".join(str(a) for a in sorted(r["anos"])), ", ".join(r["modos"]), ", ".join(r["ufs"]),
            r["declarado"], r["compensar"], r.get("regra_preco"), r.get("preco_t"), r.get("classe_mo"),
            r.get("custo_direto"), r.get("resultado"), r.get("margem"),
            r["subtotal"], r["desconto_percentual"], r["total"],
            negociado, negociado if negociado is not None else r["total"], reg.get("status"),
            _data(reg.get("data_venda")), reg.get("observacao"),
            d.get("criada_em", "").replace("T", " "), d.get("atualizada_em", "").replace("T", " "),
        ])
    _aba(ws, colunas, linhas)

    ws = wb.create_sheet("Detalhe por ano e material")
    colunas = [
        ("Nº da proposta", 22, None), ("Cliente", 34, None), ("Status", 15, None),
        ("Ano-base", 10, None), ("Modo", 18, None), ("Recuperação", 12, PCT),
        ("Material", 18, None), ("Categoria", 12, None), ("Declarado (t)", 14, TON),
        ("A compensar (t)", 14, TON), ("Valor (R$)", 16, MOEDA), ("Regra de preço", 18, None),
    ]
    linhas = [[d["numero"], d["resumo"]["cliente"], d["registro"].get("status"), x["ano"], x["modo"],
               x["recuperacao"], x["material"], x["categoria"], x["declarado"], x["compensar"],
               x["valor"], x["regra_preco"]]
              for d in propostas for x in d["resumo"].get("detalhes", [])]
    _aba(ws, colunas, linhas)

    ws = wb.create_sheet("Resumo")
    colunas = [("Status", 18, None), ("Propostas", 12, None), ("Massa a compensar (t)", 16, TON),
               ("Valor final (R$)", 18, MOEDA)]
    linhas = []
    for status in STATUS:
        grupo = [d for d in propostas if d["registro"].get("status") == status]
        valor = sum(d["registro"]["valor_negociado"] if d["registro"].get("valor_negociado") is not None
                    else d["resumo"]["total"] for d in grupo)
        linhas.append([status, len(grupo), sum(d["resumo"]["compensar"] for d in grupo), valor])
    linhas.append(["Total", len(propostas), sum(l[2] for l in linhas), sum(l[3] for l in linhas)])
    _aba(ws, colunas, linhas)
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)
    ws.freeze_panes = None

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def _atualizar_planilha():
    """Regrava a planilha do registro. Devolve um aviso se não foi possível (ex.: aberta no Excel)."""
    try:
        DADOS_DIR.mkdir(parents=True, exist_ok=True)
        conteudo = gerar_relatorio().getvalue()
        with open(PLANILHA_PATH, "wb") as f:
            f.write(conteudo)
        return None
    except PermissionError:
        return ("A proposta foi salva, mas a planilha registro_propostas.xlsx está aberta no Excel e "
                "não foi atualizada. Feche o arquivo; ela será atualizada no próximo salvamento.")
