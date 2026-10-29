"""
Cria um modelo PowerPoint PROVISÓRIO em modelo/Proposta_Comercial.pptx
só para testar o botão "Atualizar proposta" enquanto o modelo oficial
não fica pronto. Substitua o arquivo pelo PPT real com os mesmos marcadores.

Uso:  python scripts/criar_modelo_ppt_exemplo.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

DESTINO = Path(__file__).resolve().parent.parent / "modelo" / "Proposta_Comercial.pptx"

FOREST = RGBColor(0x16, 0x33, 0x00)
LIME = RGBColor(0x9F, 0xE8, 0x70)
MIST = RGBColor(0xE2, 0xF6, 0xD5)
INK = RGBColor(0x0E, 0x0F, 0x0C)
SLATE = RGBColor(0x6A, 0x6C, 0x6A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def fundo(slide, cor):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = cor


def texto(slide, x, y, w, h, linhas):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (conteudo, tamanho, cor, negrito) in enumerate(linhas):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = conteudo
        run.font.size = Pt(tamanho)
        run.font.color.rgb = cor
        run.font.bold = negrito
        run.font.name = "Inter"
        p.space_after = Pt(6)
    return box


def cartao(slide, x, y, w, h, rotulo, valor, escuro=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = FOREST if escuro else MIST
    shape.line.fill.background()
    shape.adjustments[0] = 0.12
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = rotulo
    r.font.size = Pt(12)
    r.font.color.rgb = MIST if escuro else SLATE
    p2 = tf.add_paragraph()
    r2 = p2.add_run()
    r2.text = valor
    r2.font.size = Pt(24)
    r2.font.bold = True
    r2.font.color.rgb = LIME if escuro else INK


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    branco = prs.slide_layouts[6]

    capa = prs.slides.add_slide(branco)
    fundo(capa, FOREST)
    texto(capa, 0.8, 0.6, 6, 0.5, [("ambipar", 28, LIME, True)])
    texto(capa, 0.8, 2.4, 11.5, 3, [
        ("Proposta Comercial", 54, WHITE, True),
        ("Logística reversa de embalagens", 24, MIST, False),
        ("Preparada para {{CLIENTE}}", 22, LIME, True),
    ])
    texto(capa, 0.8, 6.2, 11.5, 0.8, [
        ("Proposta nº {{NUMERO_PROPOSTA}} · {{DATA_PROPOSTA}} · válida até {{VALIDADE}}", 14, MIST, False),
    ])

    escopo = prs.slides.add_slide(branco)
    texto(escopo, 0.8, 0.6, 11.5, 1, [("Escopo da compensação", 36, INK, True)])
    texto(escopo, 0.8, 1.5, 11.5, 1.6, [
        ("Cliente: {{CLIENTE}} · CNPJ {{CNPJ}}", 16, SLATE, False),
        ("Anos de compensação: {{ANOS}}", 16, SLATE, False),
        ("Estados ({{QTD_ESTADOS}}): {{ESTADOS}}", 16, SLATE, False),
        ("Materiais: {{MATERIAIS}}", 16, SLATE, False),
    ])
    cartao(escopo, 0.8, 4.2, 3.7, 1.6, "Massa declarada", "{{VOLUME_DECLARADO}}")
    cartao(escopo, 4.8, 4.2, 3.7, 1.6, "Massa a compensar", "{{VOLUME_COMPENSAR}}")
    cartao(escopo, 8.8, 4.2, 3.7, 1.6, "Investimento", "{{VALOR_TOTAL}}", escuro=True)

    contato = prs.slides.add_slide(branco)
    texto(contato, 0.8, 0.6, 11.5, 1, [("Próximos passos", 36, INK, True)])
    texto(contato, 0.8, 1.6, 11.5, 3, [
        ("A proposta técnica com o detalhamento por ano, estado e material acompanha esta apresentação.", 16, SLATE, False),
        ("Contato no cliente: {{CONTATO}} · {{EMAIL_CLIENTE}} · {{TELEFONE}}", 16, SLATE, False),
        ("Responsável Ambipar: {{RESPONSAVEL}} · {{EMAIL_RESPONSAVEL}}", 16, SLATE, False),
        ("{{OBSERVACOES}}", 14, SLATE, False),
    ])
    texto(contato, 0.8, 6.6, 11.5, 0.5, [("{{EMPRESA}} · CNPJ {{CNPJ_EMPRESA}}", 11, SLATE, False)])

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    prs.save(DESTINO)
    print(f"Modelo provisório criado em {DESTINO}")


if __name__ == "__main__":
    main()
