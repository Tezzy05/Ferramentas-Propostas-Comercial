from flask import Flask, render_template, request, send_file, jsonify
from openpyxl import load_workbook
from io import BytesIO
from pathlib import Path
import math
import re

app = Flask(__name__)

BASE_DIR = Path(__file__).resolve().parent
MODEL_PATH = BASE_DIR / "modelo" / "Declaracao_volume_embalagens.xlsx"
OUTPUT_FILENAME = "Autodeclaracao_volume_embalagens_preenchida.xlsx"

MATERIAIS = [
    "PAPEL/PAPELÃO", "PLÁSTICO", "VIDRO", "AÇO/FERRO", "ALUMÍNIO",
    "LONGA VIDA", "ISOPOR", "BOPP", "AEROSSOL", "OUTRO MATERIAL"
]

ESTADOS = [
    "AC","AL","AM","AP","BA","CE","DF","ES","GO","MA","MG","MS","MT",
    "PA","PB","PE","PI","PR","RJ","RN","RO","RR","RS","SC","SE","SP","TO"
]

FIRST_STATE_ROW = 5
LAST_STATE_ROW = 31
FIRST_MATERIAL_COL = 2
LAST_MATERIAL_COL = 11


def load_model():
    if not MODEL_PATH.exists():
        raise FileNotFoundError(f"Modelo não encontrado: {MODEL_PATH}")
    return load_workbook(MODEL_PATH, data_only=False, read_only=False)


def get_available_years():
    wb = load_model()
    return sorted(
        [int(s) for s in wb.sheetnames if re.fullmatch(r"\d{4}", str(s))],
        reverse=True
    )


def normalize_number(value):
    if value is None or value == "":
        return 0.0

    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(" ", "")
        if "," in text and "." in text:
            text = text.replace(".", "").replace(",", ".")
        elif "," in text:
            text = text.replace(",", ".")
        try:
            number = float(text)
        except ValueError:
            raise ValueError(f"Valor inválido: {value}")

    if not math.isfinite(number):
        raise ValueError(f"Valor numérico inválido: {value}")
    if number < 0:
        raise ValueError("Não são permitidos valores negativos.")
    return number


def get_state_rows(ws):
    rows = {}
    for row in range(FIRST_STATE_ROW, LAST_STATE_ROW + 1):
        value = ws.cell(row=row, column=1).value
        if value is not None:
            state = str(value).strip().upper()
            if re.fullmatch(r"[A-Z]{2}", state):
                rows[state] = row
    return rows


def validate_payload(payload):
    if not isinstance(payload, dict):
        raise ValueError("Dados recebidos em formato inválido.")

    years_data = payload.get("anos")
    if not isinstance(years_data, dict) or not years_data:
        raise ValueError("Adicione pelo menos um estado antes de gerar o Excel.")

    wb = load_model()
    normalized = {}

    for year_text, states_data in years_data.items():
        try:
            year = int(year_text)
        except (TypeError, ValueError):
            raise ValueError(f"Ano inválido: {year_text}")

        sheet_name = str(year)
        if sheet_name not in wb.sheetnames:
            raise ValueError(f'A guia "{sheet_name}" não existe no modelo.')

        if not isinstance(states_data, dict):
            raise ValueError(f"Os dados da guia {year} estão inválidos.")

        state_rows = get_state_rows(wb[sheet_name])
        normalized[year] = {}

        for state_text, values in states_data.items():
            state = str(state_text).strip().upper()

            if state not in state_rows:
                raise ValueError(f'O estado "{state}" não foi encontrado na guia {year}.')

            if not isinstance(values, list) or len(values) != len(MATERIAIS):
                raise ValueError(f"{state}/{year}: são esperados {len(MATERIAIS)} valores.")

            vals = [normalize_number(v) for v in values]

            if all(v == 0 for v in vals):
                raise ValueError(f"{state}/{year}: informe pelo menos um material.")

            normalized[year][state] = vals

    return normalized


def generate_excel(client_data):
    """
    Usa o Excel original como template.
    NÃO recria nem formata a planilha.
    Somente preenche B5:K31.
    Todos os estados ficam com zero; os informados sobrescrevem os zeros.
    """
    wb = load_model()

    # Zera todos os campos de volume em todas as abas existentes.
    for ws in wb.worksheets:
        state_rows = get_state_rows(ws)

        for state in ESTADOS:
            if state not in state_rows:
                raise ValueError(
                    f'O estado "{state}" não foi encontrado na guia "{ws.title}".'
                )

            row = state_rows[state]

            for col in range(FIRST_MATERIAL_COL, LAST_MATERIAL_COL + 1):
                ws.cell(row=row, column=col).value = 0

    # Sobrescreve somente os valores informados.
    for year, states_data in client_data.items():
        ws = wb[str(year)]
        state_rows = get_state_rows(ws)

        for state, values in states_data.items():
            row = state_rows[state]

            for offset, value in enumerate(values):
                col = FIRST_MATERIAL_COL + offset
                ws.cell(row=row, column=col).value = value

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output


@app.route("/")
def index():
    try:
        return render_template(
            "index.html",
            anos=get_available_years(),
            materiais=MATERIAIS,
            estados=ESTADOS
        )
    except Exception as exc:
        app.logger.exception("Erro ao carregar o modelo.")
        return f"<h1>Erro ao carregar o formulário</h1><p>{exc}</p>", 500


@app.route("/gerar-excel", methods=["POST"])
def generate_excel_route():
    try:
        payload = request.get_json(silent=True)
        if payload is None:
            return jsonify({"sucesso": False, "erro": "Nenhum dado foi recebido."}), 400

        client_data = validate_payload(payload)
        output = generate_excel(client_data)

        return send_file(
            output,
            as_attachment=True,
            download_name=OUTPUT_FILENAME,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except (ValueError, FileNotFoundError) as exc:
        return jsonify({"sucesso": False, "erro": str(exc)}), 400

    except Exception:
        app.logger.exception("Erro inesperado ao gerar Excel.")
        return jsonify({
            "sucesso": False,
            "erro": "Ocorreu um erro inesperado no servidor ao gerar o Excel."
        }), 500


if __name__ == "__main__":
    print("=" * 60)
    print("AUTODECLARAÇÃO DE MASSAS - AMBIPAR")
    print("=" * 60)
    print(f"Modelo: {MODEL_PATH}")
    print("Acesse: http://127.0.0.1:5000")
    print("Para encerrar: CTRL+C")
    print("=" * 60)
    app.run(host="127.0.0.1", port=5000, debug=True)
