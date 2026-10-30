"""
Tabela "Compensação – Meta Mínima" do slide de proposta comercial.

No PPT oficial essa tabela é uma planilha do Excel embutida (objeto OLE), que o
PowerPoint mostra por uma imagem de visualização e que não se atualiza sozinha.
Aqui ela é trocada por uma tabela nativa do PowerPoint, na mesma posição e
tamanho, com o mesmo visual, preenchida com os dados de cada proposta.

Onde a tabela entra: no lugar do objeto do Excel embutido ou de uma caixa de
texto com o marcador {{TABELA_RESUMO}}.
"""
from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

from .calculos import MATERIAIS, fmt_moeda, fmt_num, fmt_pct

AZUL = RGBColor(0x00, 0xB0, 0xF0)
LIMA = RGBColor(0xD4, 0xFF, 0x00)
CINZA_CLARO = RGBColor(0xF2, 0xF2, 0xF2)
CINZA = RGBColor(0xD9, 0xD9, 0xD9)
PRETO = RGBColor(0x00, 0x00, 0x00)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
FONTE = "Neo Sans Std"
MARCADOR = "{{TABELA_RESUMO}}"
R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

# Linhas fixas como no modelo; as demais só aparecem quando têm volume.
LINHAS_FIXAS = [
    ("Papel/Papelão", ["PAPEL/PAPELÃO"]),
    ("Plástico", ["PLÁSTICO"]),
    ("Metal", ["AÇO/FERRO", "ALUMÍNIO"]),
    ("Vidro", ["VIDRO"]),
    ("Longa Vida", ["LONGA VIDA"]),
]
LINHAS_EXTRAS = [
    ("Isopor", ["ISOPOR"]),
    ("BOPP", ["BOPP"]),
    ("Aerossol", ["AEROSSOL"]),
    ("Outros", ["OUTRO MATERIAL"]),
]


def montar_tabela_resumo(proposta, calc):
    """Dados da tabela a partir do cálculo da proposta."""
    declarado = dict.fromkeys(MATERIAIS, 0.0)
    compensar = dict.fromkeys(MATERIAIS, 0.0)
    for ano in calc["por_ano"]:
        for m in ano["materiais"]:
            i = proposta["materiais"].index(m["material"])
            declarado[MATERIAIS[i]] += m["declarado"]
            compensar[MATERIAIS[i]] += m["compensar"]

    linhas = []
    for nome, chaves in LINHAS_FIXAS + LINHAS_EXTRAS:
        total = sum(declarado[c] for c in chaves)
        minimo = sum(compensar[c] for c in chaves)
        if (nome, chaves) in LINHAS_FIXAS or total > 0:
            linhas.append((nome, total, minimo))

    # Só é "mínimo" se nenhum ano contratou acima da meta legal.
    so_minimo = all(
        abs(a["recuperacao"] - a["regra"]["minimo"] / 100) < 1e-9 and not a["ajustes_uf"]
        for a in calc["por_ano"]
    )
    modos = {a["modo"] for a in calc["por_ano"]}
    abrangencia = ("Nacional" if modos == {"nacional"} else
                   "por Estado" if modos == {"estados"} else "Nacional e por Estado")
    nota = ""
    if calc["desconto_valor"]:
        nota = f'com desconto de {fmt_pct(calc["desconto_percentual"])}'
    return {
        "titulo": f'Compensação {abrangencia} – {"Meta Mínima" if so_minimo else "Meta Contratada"}',
        "coluna_volume": "Volume mínimo (ton)" if so_minimo else "A compensar (ton)",
        "linhas": linhas,
        "total_declarado": calc["declarado"],
        "total_compensar": calc["compensar"],
        "valor_t": fmt_moeda(calc["preco_efetivo"]),
        "total": fmt_moeda(calc["total"]),
        "nota": nota,
    }


def _alvo(slide):
    """Objeto do Excel embutido ou caixa com {{TABELA_RESUMO}} no slide, se houver."""
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame \
                and shape.text_frame.text.strip() == MARCADOR:
            return shape
        el = shape._element
        if el.tag == qn("p:graphicFrame") and el.find(f'.//{qn("p:oleObj")}') is not None:
            prog = el.find(f'.//{qn("p:oleObj")}').get("progId", "")
            if prog.startswith("Excel"):
                return shape
    return None


def _bordas(cell, cor="FFFFFF", largura=12700):
    tc_pr = cell._tc.get_or_add_tcPr()
    for lado in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
        for antigo in tc_pr.findall(qn(lado)):
            tc_pr.remove(antigo)
        ln = etree.SubElement(tc_pr, qn(lado), w=str(largura), cap="flat", cmpd="sng", algn="ctr")
        fill = etree.SubElement(ln, qn("a:solidFill"))
        etree.SubElement(fill, qn("a:srgbClr"), val=cor)
    # a ordem exigida pelo schema: bordas antes do preenchimento
    preenchimento = tc_pr.find(qn("a:solidFill"))
    if preenchimento is not None:
        tc_pr.remove(preenchimento)
        tc_pr.append(preenchimento)


def _celula(cell, texto, fundo, tamanho, negrito=False, cor=PRETO, alinhar=PP_ALIGN.CENTER, nota=""):
    cell.fill.solid()
    cell.fill.fore_color.rgb = fundo
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = cell.margin_right = Pt(4)
    cell.margin_top = cell.margin_bottom = Pt(1)
    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = alinhar
    run = p.add_run()
    run.text = texto
    run.font.name = FONTE
    run.font.size = Pt(tamanho)
    run.font.bold = negrito
    run.font.color.rgb = cor
    if nota:
        p2 = tf.add_paragraph()
        p2.alignment = alinhar
        r2 = p2.add_run()
        r2.text = nota
        r2.font.name = FONTE
        r2.font.size = Pt(tamanho - 4)
        r2.font.color.rgb = cor
    _bordas(cell)


def inserir_tabela(prs, dados):
    """Troca o objeto do Excel (ou o marcador) pela tabela preenchida. Devolve True se inseriu."""
    for slide in prs.slides:
        alvo = _alvo(slide)
        if alvo is None:
            continue
        left, top, width, height = alvo.left, alvo.top, alvo.width, alvo.height
        el = alvo._element
        # Relacionamentos usados só pelo objeto (planilha embutida e imagem de visualização).
        rids = {v for node in el.iter() for k, v in node.attrib.items() if k.startswith(R_NS)}
        el.getparent().remove(el)
        for rid in rids:
            slide.part.drop_rel(rid)

        linhas = dados["linhas"]
        n_linhas = len(linhas) + 3  # título, cabeçalho, total
        # Mantém a altura do modelo (8 linhas); com linhas extras, elas ficam mais baixas
        # e a fonte diminui, para não invadir o texto de baixo.
        altura_linha = int(height / max(n_linhas, 8))
        reduz = 2 if n_linhas > 8 else 0
        tabela = slide.shapes.add_table(n_linhas, 5, left, top, width, altura_linha * n_linhas).table
        tabela.first_row = False
        tabela.horz_banding = False
        for coluna in tabela.columns:
            coluna.width = int(width / 5)
        for linha in tabela.rows:
            linha.height = altura_linha

        titulo = tabela.cell(0, 0)
        titulo.merge(tabela.cell(0, 4))
        _celula(titulo, dados["titulo"], AZUL, 16 - reduz, negrito=True, cor=BRANCO)

        cabecalho = ["Material", "Volume total (ton)", dados["coluna_volume"], "Valor / Ton", "Total (R$)"]
        for j, texto in enumerate(cabecalho):
            _celula(tabela.cell(1, j), texto, LIMA, 13 - reduz, negrito=True)

        for k, (nome, total, minimo) in enumerate(linhas):
            fundo = CINZA_CLARO if k % 2 == 0 else CINZA
            r = 2 + k
            _celula(tabela.cell(r, 0), nome, fundo, 13 - reduz)
            _celula(tabela.cell(r, 1), fmt_num(total, 3), fundo, 13 - reduz)
            _celula(tabela.cell(r, 2), fmt_num(minimo, 3), fundo, 13 - reduz)

        ultima = n_linhas - 1
        _celula(tabela.cell(ultima, 0), "Total", LIMA, 13 - reduz, negrito=True)
        _celula(tabela.cell(ultima, 1), fmt_num(dados["total_declarado"], 3), LIMA, 13 - reduz, negrito=True)
        _celula(tabela.cell(ultima, 2), fmt_num(dados["total_compensar"], 3), LIMA, 13 - reduz, negrito=True)

        valor = tabela.cell(2, 3)
        valor.merge(tabela.cell(ultima, 3))
        _celula(valor, dados["valor_t"], CINZA_CLARO, 15)
        total = tabela.cell(2, 4)
        total.merge(tabela.cell(ultima, 4))
        _celula(total, dados["total"], CINZA_CLARO, 15, negrito=True, nota=dados["nota"])
        return True
    return False
